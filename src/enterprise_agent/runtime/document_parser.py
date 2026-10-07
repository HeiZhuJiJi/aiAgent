import csv
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from pathlib import PurePosixPath

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from enterprise_agent.contracts.domain import FailureCode
from enterprise_agent.domain.states import DomainError

MAX_DOCUMENT_BYTES = 50 * 1024 * 1024
MAX_EXPANDED_BYTES = 200 * 1024 * 1024
SUPPORTED_DOCUMENT_MEDIA_TYPES = frozenset(
    {
        "application/pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "text/csv",
        "text/markdown",
        "text/plain",
    }
)


@dataclass(frozen=True)
class ParsedPart:
    content: str
    locator: dict[str, str | int]


def _normalize(value: str) -> str:
    value = (
        unicodedata.normalize("NFKC", value)
        .replace("\x00", "")
        .replace("\r\n", "\n")
        .replace("\r", "\n")
    )
    return re.sub(r"[ \t]+", " ", value).strip()


def _validate_office_archive(blob: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(blob)) as archive:
            total = 0
            for member in archive.infolist():
                path = PurePosixPath(member.filename)
                if path.is_absolute() or ".." in path.parts or "\\" in member.filename:
                    raise ValueError("unsafe archive path")
                total += member.file_size
                if total > MAX_EXPANDED_BYTES or member.file_size > max(member.compress_size, 1) * 1000:
                    raise ValueError("archive expansion limit")
    except (zipfile.BadZipFile, ValueError):
        raise DomainError(FailureCode.VALIDATION_FAILED, "unsafe office document") from None


def parse_document(blob: bytes, media_type: str) -> tuple[ParsedPart, ...]:
    if not blob or len(blob) > MAX_DOCUMENT_BYTES:
        raise DomainError(FailureCode.VALIDATION_FAILED, "document size invalid")
    parts: list[ParsedPart] = []
    try:
        if media_type == "application/pdf":
            for page, page_object in enumerate(PdfReader(io.BytesIO(blob)).pages, start=1):
                content = _normalize(page_object.extract_text() or "")
                if content:
                    parts.append(ParsedPart(content, {"page": page}))
        elif media_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            _validate_office_archive(blob)
            for paragraph, paragraph_object in enumerate(
                Document(io.BytesIO(blob)).paragraphs, start=1
            ):
                content = _normalize(paragraph_object.text)
                if content:
                    parts.append(
                        ParsedPart(
                            content,
                            {
                                "paragraph": paragraph,
                                "section": (
                                    paragraph_object.style.name
                                    if paragraph_object.style is not None
                                    else "Normal"
                                ),
                            },
                        )
                    )
        elif media_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
            _validate_office_archive(blob)
            for slide, value in enumerate(Presentation(io.BytesIO(blob)).slides, start=1):
                content = _normalize("\n".join(shape.text for shape in value.shapes if hasattr(shape, "text")))
                if content:
                    parts.append(ParsedPart(content, {"slide": slide}))
        elif media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            _validate_office_archive(blob)
            workbook = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
            try:
                for sheet in workbook.worksheets:
                    for row_number, sheet_row in enumerate(sheet.iter_rows(), start=1):
                        cell_values = [
                            f"{cell.coordinate}={cell.value}"
                            for cell in sheet_row
                            if cell.value is not None
                        ]
                        if cell_values:
                            parts.append(
                                ParsedPart(
                                    _normalize(" | ".join(cell_values)),
                                    {
                                        "sheet": sheet.title,
                                        "row": row_number,
                                        "cells": (
                                            f"{sheet_row[0].coordinate}:"
                                            f"{sheet_row[-1].coordinate}"
                                        ),
                                    },
                                )
                            )
            finally:
                workbook.close()
        elif media_type == "text/csv":
            for row_number, csv_row in enumerate(
                csv.reader(io.StringIO(blob.decode("utf-8-sig"))), start=1
            ):
                content = _normalize(" | ".join(csv_row))
                if content:
                    parts.append(ParsedPart(content, {"row": row_number}))
        elif media_type in {"text/markdown", "text/plain"}:
            for line_number, line in enumerate(blob.decode("utf-8").splitlines(), start=1):
                content = _normalize(line)
                if content:
                    parts.append(ParsedPart(content, {"line": line_number}))
        else:
            raise DomainError(FailureCode.VALIDATION_FAILED, "unsupported document media type")
    except DomainError:
        raise
    except Exception as exc:  # noqa: BLE001 - document parsers are an external trust boundary
        raise DomainError(FailureCode.VALIDATION_FAILED, f"document parse failed: {type(exc).__name__}") from None
    if not parts:
        raise DomainError(FailureCode.VALIDATION_FAILED, "document has no extractable content")
    return tuple(parts)
