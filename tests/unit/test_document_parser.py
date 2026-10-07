import io
import zipfile

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from enterprise_agent.domain.states import DomainError
from enterprise_agent.runtime.document_parser import parse_document


def pdf_bytes() -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 72 720 Td (Supplier risk) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def test_pdf_docx_pptx_and_xlsx_locators() -> None:
    pdf = parse_document(pdf_bytes(), "application/pdf")
    assert pdf[0].content == "Supplier risk" and pdf[0].locator == {"page": 1}

    document = Document()
    document.add_heading("Supplier Profile", level=1)
    document.add_paragraph("Certified supplier")
    value = io.BytesIO()
    document.save(value)
    docx = parse_document(value.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert docx[0].locator["paragraph"] == 1 and docx[1].content == "Certified supplier"

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Risk Rating"
    slide.placeholders[1].text = "LOW"
    value = io.BytesIO()
    presentation.save(value)
    pptx = parse_document(value.getvalue(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")
    assert pptx[0].locator == {"slide": 1} and "Risk Rating" in pptx[0].content

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Risk"
    sheet.append(["Supplier", "Rating"])
    sheet.append(["ACME", "LOW"])
    value = io.BytesIO()
    workbook.save(value)
    xlsx = parse_document(value.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert xlsx[1].locator == {"sheet": "Risk", "row": 2, "cells": "A2:B2"}


@pytest.mark.parametrize(
    ("media_type", "blob", "locator"),
    [
        ("text/csv", b"supplier,risk\nACME,LOW\n", {"row": 1}),
        ("text/markdown", b"# Report\n\nEvidence\n", {"line": 1}),
        ("text/plain", b"Supplier profile\n", {"line": 1}),
    ],
)
def test_text_formats_normalize_and_locate(media_type, blob, locator) -> None:
    parts = parse_document(blob, media_type)
    assert parts[0].locator == locator


def test_database_forbidden_nul_bytes_are_removed_at_parser_boundary() -> None:
    parts = parse_document(b"safe\x00 text\n", "text/plain")

    assert parts[0].content == "safe text"
    assert "\x00" not in parts[0].content


def test_parser_rejects_unsupported_and_expanding_office_archives() -> None:
    with pytest.raises(DomainError, match="unsupported"):
        parse_document(b"data", "application/octet-stream")
    value = io.BytesIO()
    with zipfile.ZipFile(value, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"0" * 2_000_000)
    with pytest.raises(DomainError, match="unsafe office"):
        parse_document(value.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
