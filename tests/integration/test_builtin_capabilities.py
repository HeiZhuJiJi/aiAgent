import hashlib
import io
import json
import os
import threading
from datetime import UTC, datetime, timedelta
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from typing import ClassVar
from uuid import UUID, uuid4

from docx import Document
from dotenv import dotenv_values
from jsonschema import Draft202012Validator
from minio import Minio
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from enterprise_agent.capabilities.builtins import (
    builtin_descriptors,
    register_builtin_capabilities,
)
from enterprise_agent.capabilities.registry import CapabilityRegistry
from enterprise_agent.contracts.domain import CapabilityRequest, Evidence, TaskGoal
from enterprise_agent.runtime.budgets import new_task_budget
from enterprise_agent.storage.artifacts import ArtifactStore
from enterprise_agent.storage.knowledge import ChunkInput, persist_document_version
from enterprise_agent.storage.plans import persist_snapshot
from enterprise_agent.storage.provenance import load_claims
from enterprise_agent.storage.tables import (
    AttemptRow,
    EvidenceRow,
    RuntimeGuardRow,
    StepRow,
    TaskBudgetRow,
    TaskRow,
    Workspace,
)
from enterprise_agent.storage.tasks import current_snapshot


class BuiltinProviderHandler(BaseHTTPRequestHandler):
    calls: ClassVar[dict[str, int]] = {}

    def do_POST(self) -> None:
        length = int(self.headers["Content-Length"])
        request = json.loads(self.rfile.read(length))
        type(self).calls[self.path] = type(self).calls.get(self.path, 0) + 1
        if self.path == "/v1/embeddings":
            vector = [0.0] * 1024
            vector[0] = 1
            payload = {"data": [{"index": 0, "embedding": vector}]}
        elif self.path == "/v1/rerank":
            payload = {
                "results": [
                    {
                        "index": index,
                        "relevance_score": (
                            0.95 if "sanctions" in content.casefold() else 0.1
                        ),
                    }
                    for index, content in enumerate(request["documents"])
                ]
            }
        elif self.path == "/v1/chat/completions":
            evidence_payload = json.loads(request["messages"][1]["content"])
            evidence_id = evidence_payload["evidence"][0]["evidence_id"]
            output = {
                "schema_version": 1,
                "decision": "ADMIT",
                "risk_level": "LOW",
                "summary": "The cited screening result is clear.",
                "key_findings": ["No sanctions match was found."],
                "missing_information": [],
                "recommendation": "Admit after human approval.",
                "evidence_sufficient": True,
                "claims": [
                    {
                        "schema_version": 1,
                        "content": "The sanctions screening is clear.",
                        "importance": "HIGH",
                        "risk_level": "READ",
                        "evidence_ids": [evidence_id],
                    }
                ],
            }
            payload = {
                "model": "builtin-test-model",
                "choices": [{"message": {"content": json.dumps(output)}}],
                "usage": {"prompt_tokens": 30, "completion_tokens": 20},
            }
        elif self.path == "/v1/search":
            snippet = "Public registry and sanctions search returned no adverse match."
            payload = {
                "results": [
                    {
                        "schema_version": 1,
                        "title": "Public supplier record",
                        "url": "https://example.test/supplier",
                        "snippet": snippet,
                        "published_at": None,
                        "observed_at": "2026-09-12T00:00:00Z",
                        "content_hash": hashlib.sha256(snippet.encode()).hexdigest(),
                    }
                ]
            }
        elif self.path == "/v1/fetch":
            content = "Verified public supplier page."
            payload = {
                "schema_version": 1,
                "url": request["url"],
                "content": content,
                "content_hash": hashlib.sha256(content.encode()).hexdigest(),
            }
        else:
            self.send_error(404)
            return
        body = json.dumps(payload).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, format: str, *args: object) -> None:
        return


def _request(
    *,
    workspace_id,
    task_id,
    step_id,
    attempt_id,
    capability_id,
    arguments,
):
    return CapabilityRequest(
        command_id=uuid4(),
        operation_id=uuid4(),
        workspace_id=workspace_id,
        task_id=task_id,
        step_id=step_id,
        attempt_id=attempt_id,
        fencing_token=1,
        capability_id=capability_id,
        capability_version="1.0.0",
        arguments=arguments,
    )


async def test_builtin_capabilities_use_real_postgres_http_model_web_and_minio(
    monkeypatch,
):
    password = dotenv_values(".env").get("MINIO_ROOT_PASSWORD")
    assert password
    minio_endpoint = os.environ.get("TEST_MINIO_ENDPOINT", "127.0.0.1:59000")
    server = ThreadingHTTPServer(("127.0.0.1", 0), BuiltinProviderHandler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    base_url = f"http://127.0.0.1:{server.server_port}/v1"
    for prefix, name in (
        ("MODEL", "builtin-test-model"),
        ("EMBEDDING", "BAAI/bge-m3"),
        ("RERANKER", "BAAI/bge-reranker-v2-m3"),
    ):
        monkeypatch.setenv(prefix + "_MODE", "local")
        monkeypatch.setenv(prefix + "_BASE_URL", base_url)
        monkeypatch.setenv(prefix + "_NAME", name)
        monkeypatch.setenv(prefix + "_API_KEY", "")
    monkeypatch.setenv("WEB_MODE", "local")
    monkeypatch.setenv("WEB_BASE_URL", base_url)
    monkeypatch.setenv("WEB_API_KEY", "")
    monkeypatch.setenv("MINIO_ENDPOINT", minio_endpoint)
    monkeypatch.setenv("MINIO_ACCESS_KEY", "platform")
    monkeypatch.setenv("MINIO_SECRET_KEY", str(password))
    bucket = "runtime-builtins-" + uuid4().hex
    monkeypatch.setenv("ARTIFACT_BUCKET", bucket)

    engine = create_async_engine(os.environ["TEST_DATABASE_URL"])
    sessions = async_sessionmaker(engine, expire_on_commit=False)
    workspace_id, task_id, step_id, attempt_id = (
        uuid4(),
        uuid4(),
        uuid4(),
        uuid4(),
    )
    goal = TaskGoal(
        goal="Assess the supplier",
        constraints=("Cite Evidence",),
        deliverables=("Supplier Risk Report",),
        workspace_id=workspace_id,
        time_budget_sec=600,
        token_budget=10_000,
        cost_budget=10,
    )
    source_evidence = Evidence(
        evidence_id=uuid4(),
        workspace_id=workspace_id,
        task_id=task_id,
        step_id=step_id,
        source_type="MCP",
        source_id="procurement.get_supplier_risk_flags",
        locator="mcp:risk-flags",
        content="Sanctions screening status: CLEAR.",
        content_hash=hashlib.sha256(b"Sanctions screening status: CLEAR.").hexdigest(),
        source_hash=hashlib.sha256(b"Sanctions screening status: CLEAR.").hexdigest(),
        source_version="1.0.0",
        confidence=1,
        relevance=1,
        retrieved_at=datetime.now(UTC),
        generated_by=attempt_id,
        derived_from=(),
    )
    execution_snapshot = current_snapshot(workspace_id)
    async with sessions.begin() as session:
        session.add(Workspace(id=workspace_id, name="builtin-capabilities"))
        await session.flush()
        await persist_snapshot(session, execution_snapshot)
        session.add(
            TaskRow(
                id=task_id,
                workspace_id=workspace_id,
                state="RUNNING",
                execution_snapshot_id=execution_snapshot.snapshot_id,
                goal_payload=goal.model_dump(mode="json"),
            )
        )
        await session.flush()
        session.add(new_task_budget(task_id, goal))
        session.add(RuntimeGuardRow(task_id=task_id, workspace_id=workspace_id))
        session.add(
            StepRow(
                id=step_id,
                workspace_id=workspace_id,
                task_id=task_id,
                state="RUNNING",
                fencing_token=1,
                lease_expires_at=datetime.now(UTC) + timedelta(minutes=5),
            )
        )
        await session.flush()
        session.add(
            AttemptRow(
                id=attempt_id,
                workspace_id=workspace_id,
                step_id=step_id,
                worker_id="builtin-test",
                fencing_token=1,
                state="RUNNING",
                lease_expires_at=datetime.now(UTC) + timedelta(minutes=5),
                trace_id="trace-builtins",
            )
        )
        await session.flush()
        session.add(
            EvidenceRow(
                id=source_evidence.evidence_id,
                workspace_id=workspace_id,
                step_id=step_id,
                attempt_id=attempt_id,
                payload=source_evidence.model_dump(mode="json"),
            )
        )
        await persist_document_version(
            session,
            workspace_id=workspace_id,
            document_id=uuid4(),
            name="policy.md",
            media_type="text/markdown",
            source_hash="c" * 64,
            storage_uri="s3://knowledge/policy.md",
            parser_version="test",
            chunks=(
                ChunkInput(
                    content="Supplier sanctions screening must be clear.",
                    locator={"line": 1},
                    token_count=7,
                    embedding=(1.0,) + (0.0,) * 1023,
                ),
                ChunkInput(
                    content="Payment terms are net thirty days.",
                    locator={"line": 2},
                    token_count=7,
                    embedding=(0.0,) * 1024,
                ),
            ),
        )

    registry = CapabilityRegistry()
    register_builtin_capabilities(registry, sessions)
    descriptors = {item.id: item for item in builtin_descriptors()}
    assert set(descriptors) == {
        "knowledge.hybrid_search",
        "web.search",
        "web.fetch",
        "model.structured_output",
        "runtime.review",
        "artifact.docx",
        "artifact.markdown",
    }
    try:
        knowledge = await registry.resolve(
            "knowledge.hybrid_search", "1.0.0"
        ).handler(
            _request(
                workspace_id=workspace_id,
                task_id=task_id,
                step_id=step_id,
                attempt_id=attempt_id,
                capability_id="knowledge.hybrid_search",
                arguments={"query": "sanctions", "limit": 2},
            )
        )
        assert knowledge.outcome == "SUCCEEDED"
        Draft202012Validator(descriptors["knowledge.hybrid_search"].output_schema).validate(
            knowledge.output
        )
        assert knowledge.output["hits"][0]["locator"] == '{"line": 1}'

        web = await registry.resolve("web.search", "1.0.0").handler(
            _request(
                workspace_id=workspace_id,
                task_id=task_id,
                step_id=step_id,
                attempt_id=attempt_id,
                capability_id="web.search",
                arguments={"query": "supplier sanctions", "limit": 3},
            )
        )
        assert web.outcome == "SUCCEEDED"
        Draft202012Validator(descriptors["web.search"].output_schema).validate(web.output)

        model = await registry.resolve("model.structured_output", "1.0.0").handler(
            _request(
                workspace_id=workspace_id,
                task_id=task_id,
                step_id=step_id,
                attempt_id=attempt_id,
                capability_id="model.structured_output",
                arguments={"instruction": "Assess supplier risk."},
            )
        )
        assert model.outcome == "SUCCEEDED"
        Draft202012Validator(descriptors["model.structured_output"].output_schema).validate(
            model.output
        )
        async with sessions() as session:
            claims = await load_claims(session, workspace_id, task_id)
        assert len(claims) == 1
        assert claims[0].evidence_ids == (source_evidence.evidence_id,)

        review = await registry.resolve("runtime.review", "1.0.0").handler(
            _request(
                workspace_id=workspace_id,
                task_id=task_id,
                step_id=step_id,
                attempt_id=attempt_id,
                capability_id="runtime.review",
                arguments={"instruction": "Review Claims."},
            )
        )
        assert review.outcome == "SUCCEEDED"
        Draft202012Validator(descriptors["runtime.review"].output_schema).validate(
            review.output
        )
        assert review.output["accepted"] is True

        artifact = await registry.resolve("artifact.docx", "1.0.0").handler(
            _request(
                workspace_id=workspace_id,
                task_id=task_id,
                step_id=step_id,
                attempt_id=attempt_id,
                capability_id="artifact.docx",
                arguments={"instruction": "Create the final report."},
            )
        )
        assert artifact.outcome == "SUCCEEDED"
        Draft202012Validator(descriptors["artifact.docx"].output_schema).validate(
            artifact.output
        )
        async with sessions() as session:
            content = await ArtifactStore(
                Minio(
                    minio_endpoint,
                    access_key="platform",
                    secret_key=str(password),
                    secure=False,
                ),
                bucket=bucket,
            ).get(
                session,
                workspace_id,
                UUID(artifact.output["artifact_id"]),
            )
        report = Document(io.BytesIO(content))
        headings = {paragraph.text for paragraph in report.paragraphs}
        assert {
            "Executive Summary",
            "Supplier Profile",
            "Risk Rating",
            "Risk Matrix",
            "Key Findings",
            "Evidence",
            "Missing Information",
            "Recommendation",
            "Citation",
        } <= headings
        async with sessions() as session:
            budget = await session.get(TaskBudgetRow, task_id)
        assert budget is not None
        assert budget.used_llm_calls == 3
        assert budget.used_tokens == 50
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)
        await engine.dispose()
