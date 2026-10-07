import asyncio
import hashlib
import json
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import ClassVar
from uuid import UUID, uuid4

import httpx
import pytest
from docx import Document
from dotenv import dotenv_values
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from temporalio.client import Client
from temporalio.service import RPCError, RPCStatusCode

from apps.api.main import create_app
from enterprise_agent.procurement.demo import seed_demo_suppliers
from enterprise_agent.procurement.service import AdmissionRow
from enterprise_agent.runtime.outbox_dispatcher import dispatch_once
from enterprise_agent.storage.knowledge import ChunkInput, persist_document_version
from enterprise_agent.storage.tables import (
    ArtifactRow,
    ClaimRow,
    ReplanDecisionRow,
)
from tests.failure.test_worker_crash import start_worker
from tests.integration.test_api_identity import register_and_login


class FlagshipProvider(BaseHTTPRequestHandler):
    calls: ClassVar[dict[str, int]] = {}
    decision: ClassVar[str] = "ADMIT"
    supplier_label: ClassVar[str] = "Supplier A"
    force_replan: ClassVar[bool] = False
    gap_calls: ClassVar[int] = 0

    def do_POST(self) -> None:
        request = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        type(self).calls[self.path] = type(self).calls.get(self.path, 0) + 1
        if self.path == "/v1/embeddings":
            vector = [0.0] * 1024
            vector[0] = 1
            payload = {
                "data": [
                    {"index": index, "embedding": vector}
                    for index, _ in enumerate(request["input"])
                ]
            }
        elif self.path == "/v1/rerank":
            payload = {
                "results": [
                    {"index": index, "relevance_score": 0.99 - index * 0.01}
                    for index, _ in enumerate(request["documents"])
                ]
            }
        elif self.path == "/v1/search":
            snippet = (
                "Public registry is active and sanctions screening returned no adverse match."
            )
            payload = {
                "results": [
                    {
                        "schema_version": 1,
                        "title": type(self).supplier_label + " public screening",
                        "url": "https://example.test/suppliers/a",
                        "snippet": snippet,
                        "published_at": None,
                        "observed_at": "2026-09-12T00:00:00Z",
                        "content_hash": hashlib.sha256(snippet.encode()).hexdigest(),
                    }
                ]
            }
        elif self.path == "/v1/chat/completions":
            supplied = json.loads(request["messages"][1]["content"])
            if "deterministic_trigger" in supplied:
                output = {
                    "schema_version": 1,
                    "reason": "Collect one additional public record before deciding.",
                    "trigger": "EVIDENCE_INSUFFICIENT",
                    "action": "EXPAND_STEP",
                    "affected_steps": ["detect-evidence-gaps"],
                    "new_steps": [
                        {
                            "schema_version": 1,
                            "step_key": "collect-additional-public-records",
                            "objective": "Retrieve an additional public record for Supplier C.",
                            "dependencies": ["analyze-risk"],
                            "capability_requirements": [
                                {
                                    "schema_version": 1,
                                    "capability_id": "web.search",
                                    "version_range": ">=1,<2",
                                }
                            ],
                            "expected_output": "Additional public Evidence",
                            "success_criteria": ["A stable URL and content hash are retained"],
                            "risk_hint": "READ",
                            "execution_mode": "CAPABILITY",
                            "parallelizable": False,
                            "max_attempts": 2,
                        },
                        {
                            "schema_version": 1,
                            "step_key": "detect-evidence-gaps",
                            "objective": "Decide whether missing evidence remains after Replan.",
                            "dependencies": ["collect-additional-public-records"],
                            "capability_requirements": [
                                {
                                    "schema_version": 1,
                                    "capability_id": "model.structured_output",
                                    "version_range": ">=1,<2",
                                }
                            ],
                            "expected_output": "Resolved evidence gap decision",
                            "success_criteria": ["Unknown facts remain unknown"],
                            "risk_hint": "READ",
                            "execution_mode": "CAPABILITY",
                            "parallelizable": False,
                            "max_attempts": 2,
                        },
                    ],
                    "preserve_completed_steps": True,
                    "budget_impact": 0,
                }
                body = json.dumps(
                    {
                        "model": "flagship-local-protocol-model",
                        "choices": [{"message": {"content": json.dumps(output)}}],
                        "usage": {"prompt_tokens": 80, "completion_tokens": 40},
                    }
                ).encode()
                self.send_response(200)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)
                return
            evidence = supplied["evidence"]
            instruction = supplied["instruction"]
            gap_detection = "evidence" in instruction.casefold() and "decide" in instruction.casefold()
            insufficient = False
            if type(self).force_replan and gap_detection:
                type(self).gap_calls += 1
                insufficient = type(self).gap_calls == 1
            cited = next(
                (item for item in evidence if item["source_type"] != "DERIVED"),
                evidence[0],
            )
            output = {
                "schema_version": 1,
                "decision": type(self).decision,
                "risk_level": "LOW" if type(self).decision == "ADMIT" else "HIGH",
                "summary": (
                    type(self).supplier_label
                    + " was assessed from the cited admission requirements."
                ),
                "key_findings": [
                    "The supplier status and screening outcome are supported by cited evidence."
                ],
                "missing_information": [],
                "recommendation": (
                    "Admit after the required human approval."
                    if type(self).decision == "ADMIT"
                    else "Do not admit this supplier."
                ),
                "evidence_sufficient": not insufficient,
                "claims": [
                    {
                        "schema_version": 1,
                        "content": (
                            type(self).supplier_label
                            + " was evaluated against the cited screening."
                        ),
                        "importance": "HIGH",
                        "risk_level": "READ",
                        "evidence_ids": [cited["evidence_id"]],
                    }
                ],
            }
            payload = {
                "model": "flagship-local-protocol-model",
                "choices": [{"message": {"content": json.dumps(output)}}],
                "usage": {"prompt_tokens": 120, "completion_tokens": 45},
            }
        else:
            self.send_error(404)
            return
        body = json.dumps(payload).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, format: str, *args: object) -> None:
        return


@pytest.mark.parametrize(
    (
        "supplier_index",
        "supplier_label",
        "decision",
        "expect_approval",
        "expect_replan",
    ),
    (
        (0, "Supplier A", "ADMIT", True, False),
        (1, "Supplier B", "REJECT", False, False),
        (2, "Supplier C", "REJECT", False, True),
    ),
)
async def test_supplier_full_temporal_evidence_report_and_conditional_admission(
    monkeypatch,
    supplier_index: int,
    supplier_label: str,
    decision: str,
    expect_approval: bool,
    expect_replan: bool,
) -> None:
    password = dotenv_values(".env").get("MINIO_ROOT_PASSWORD")
    assert password
    provider = ThreadingHTTPServer(("127.0.0.1", 0), FlagshipProvider)
    provider_thread = threading.Thread(target=provider.serve_forever, daemon=True)
    provider_thread.start()
    base_url = f"http://127.0.0.1:{provider.server_port}/v1"
    for prefix, name in (
        ("MODEL", "flagship-local-protocol-model"),
        ("EMBEDDING", "BAAI/bge-m3"),
        ("RERANKER", "BAAI/bge-reranker-v2-m3"),
    ):
        monkeypatch.setenv(prefix + "_MODE", "local")
        monkeypatch.setenv(prefix + "_BASE_URL", base_url)
        monkeypatch.setenv(prefix + "_NAME", name)
        monkeypatch.setenv(prefix + "_API_KEY", "")
    monkeypatch.setenv("WEB_MODE", "local")
    monkeypatch.setenv("WEB_BASE_URL", base_url)
    monkeypatch.setenv("WEB_API_KEY", "")
    monkeypatch.setenv("MINIO_ENDPOINT", "127.0.0.1:59000")
    monkeypatch.setenv("MINIO_ACCESS_KEY", "platform")
    monkeypatch.setenv("MINIO_SECRET_KEY", str(password))
    monkeypatch.setenv("ARTIFACT_BUCKET", "flagship-artifacts-" + uuid4().hex)
    monkeypatch.setenv("SKILL_QUARANTINE_BUCKET", "flagship-skills-" + uuid4().hex)
    task_queue = "supplier-a-flagship-" + str(uuid4())
    monkeypatch.setenv("TEMPORAL_TASK_QUEUE", task_queue)

    FlagshipProvider.calls = {}
    FlagshipProvider.decision = decision
    FlagshipProvider.supplier_label = supplier_label
    FlagshipProvider.force_replan = expect_replan
    FlagshipProvider.gap_calls = 0
    app = create_app()
    temporal = await Client.connect("127.0.0.1:57233")
    procurement_engine = create_async_engine(
        str(__import__("os").environ["PROCUREMENT_DATABASE_URL"])
    )
    procurement = async_sessionmaker(procurement_engine, expire_on_commit=False)
    worker = None
    handle = None
    completed = False
    try:
        async with (
            app.router.lifespan_context(app),
            httpx.AsyncClient(
                transport=httpx.ASGITransport(app=app), base_url="http://test"
            ) as client,
        ):
            owner = await register_and_login(
                client, supplier_label.casefold().replace(" ", "-") + "-flagship"
            )
            workspace_id = UUID(
                (
                    await client.post(
                        "/workspaces",
                        headers=owner,
                        json={"name": supplier_label + " Flagship"},
                    )
                ).json()["workspace_id"]
            )
            async with procurement.begin() as session:
                supplier = (await seed_demo_suppliers(session, workspace_id))[supplier_index]

            embedding = (1.0,) + (0.0,) * 1023
            for name, content in (
                (
                    "supplier-a-submission.txt",
                    supplier_label
                    + " submitted its ownership declarations and supporting materials.",
                ),
                (
                    "procurement-policy.txt",
                    "Procurement policy permits active suppliers with clear sanctions screening after human approval.",
                ),
            ):
                source_hash = hashlib.sha256(content.encode()).hexdigest()
                async with app.state.sessions.begin() as session:
                    await persist_document_version(
                        session,
                        workspace_id=workspace_id,
                        document_id=uuid4(),
                        name=name,
                        media_type="text/plain",
                        source_hash=source_hash,
                        storage_uri="s3://flagship-source/" + source_hash,
                        parser_version="1.0.0",
                        chunks=(
                            ChunkInput(
                                content=content,
                                locator={"line": 1},
                                token_count=len(content.split()),
                                embedding=embedding,
                            ),
                        ),
                    )

            registered = (
                await client.post(
                    "/mcp",
                    headers=owner,
                    json={
                        "workspace_id": str(workspace_id),
                        "name": "Procurement",
                        "provider": "procurement",
                    },
                )
            ).json()
            connected = await client.post(
                f"/mcp/{registered['server_id']}/connect",
                headers=owner,
                json={"workspace_id": str(workspace_id), "expected_version": 1},
            )
            assert connected.status_code == 200

            skill_blob = Path(
                "demo/supplier-risk/generated/supplier-risk.skill.zip"
            ).read_bytes()
            quarantined = await client.post(
                "/skills/quarantine",
                params={
                    "workspace_id": str(workspace_id),
                    "name": "supplier-risk",
                    "version": "1.0.0",
                },
                headers={**owner, "Content-Type": "application/zip"},
                content=skill_blob,
            )
            assert quarantined.status_code == 201
            installation_id = quarantined.json()["installation_id"]
            installed = await client.post(
                "/skills/install",
                headers=owner,
                json={
                    "workspace_id": str(workspace_id),
                    "installation_id": installation_id,
                    "expected_version": 2,
                },
            )
            enabled = await client.post(
                f"/skills/{installation_id}/enable",
                headers=owner,
                json={"workspace_id": str(workspace_id), "expected_version": 3},
            )
            assert installed.status_code == enabled.status_code == 200

            stale_user_hash = "0" * 64
            created = await client.post(
                "/tasks",
                headers=owner,
                json={
                    "goal": (
                        f"Assess {supplier_label} using submitted files, procurement policy, "
                        "history and public information; produce a report and admit if eligible."
                    ),
                    "constraints": ["all material findings require citations"],
                    "deliverables": ["formal supplier risk report"],
                    "workspace_id": str(workspace_id),
                    "time_budget_sec": 900,
                    "token_budget": 50_000,
                    "cost_budget": 10,
                    "requested_skill": "supplier-risk",
                    "user_context": json.dumps(
                        {
                            "supplier_id": str(supplier.supplier_id),
                            "report_sha256": stale_user_hash,
                        }
                    ),
                },
            )
            assert created.status_code == 201, created.text
            task_id = created.json()["task_id"]
            published = await dispatch_once(
                app.state.sessions,
                temporal,
                "supplier-a-flagship-publisher",
                workspace_id,
            )
            assert published is not None
            handle = temporal.get_workflow_handle("task-" + task_id)
            worker = start_worker(task_queue)

            result = None
            bound_hash = None
            if expect_approval:
                approval = None
                for _ in range(600):
                    response = await client.get(
                        "/approvals",
                        headers=owner,
                        params={"workspace_id": str(workspace_id)},
                    )
                    if response.status_code == 200 and response.json():
                        approval = response.json()[0]
                        break
                    if worker.poll() is not None:
                        _, error = worker.communicate(timeout=1)
                        raise AssertionError(
                            "flagship Worker exited: "
                            + error.decode(errors="replace")[-4000:]
                        )
                    await asyncio.sleep(0.1)
                if approval is None:
                    step_diagnostics = (
                        await client.get(
                            f"/workspaces/{workspace_id}/tasks/{task_id}/steps",
                            headers=owner,
                        )
                    ).json()
                    attempt_diagnostics = (
                        await client.get(
                            f"/workspaces/{workspace_id}/tasks/{task_id}/attempts",
                            headers=owner,
                        )
                    ).json()
                    raise AssertionError(
                        "flagship approval deadline exceeded: "
                        + json.dumps(
                            {
                                "provider_calls": FlagshipProvider.calls,
                                "steps": step_diagnostics,
                                "attempts": attempt_diagnostics,
                            },
                            default=str,
                        )
                    )
                assert approval["state"] == "PENDING"
                bound_hash = approval["request"]["parameters"]["request"][
                    "report_sha256"
                ]
                assert bound_hash != stale_user_hash
                assert approval["request"]["evidence_ids"]
                async with procurement() as session:
                    assert (
                        await session.scalar(
                            select(func.count()).select_from(AdmissionRow).where(
                                AdmissionRow.workspace_id == workspace_id
                            )
                        )
                        == 0
                    )

                approval_id = UUID(approval["approval_id"])
                approval_handle = temporal.get_workflow_handle(
                    "approval-wait-" + str(approval_id)
                )
                for _ in range(200):
                    try:
                        if await approval_handle.query("waiting_for_approval"):
                            break
                    except RPCError as exc:
                        if exc.status != RPCStatusCode.NOT_FOUND:
                            raise
                    await asyncio.sleep(0.1)
                else:
                    raise AssertionError("flagship approval child did not become ready")
                approved = await client.post(
                    f"/approvals/{approval_id}/approve",
                    headers=owner,
                    json={"workspace_id": str(workspace_id), "expected_version": 1},
                )
                assert approved.status_code == 200
                approval_event = await dispatch_once(
                    app.state.sessions,
                    temporal,
                    "supplier-approval-publisher",
                    workspace_id,
                )
                assert approval_event is not None
            else:
                try:
                    result = await asyncio.wait_for(handle.result(), timeout=90)
                except TimeoutError:
                    if worker is not None and worker.poll() is not None:
                        _, error = worker.communicate(timeout=1)
                        raise AssertionError(
                            "flagship Worker exited: "
                            + error.decode(errors="replace")[-8000:]
                        ) from None
                    raise
                approvals = (
                    await client.get(
                        "/approvals",
                        headers=owner,
                        params={"workspace_id": str(workspace_id)},
                    )
                ).json()
                assert approvals == []

            if result is None:
                result = await asyncio.wait_for(handle.result(), timeout=90)
            assert result["state"] == "COMPLETED"
            steps = (
                await client.get(
                    f"/workspaces/{workspace_id}/tasks/{task_id}/steps",
                    headers=owner,
                )
            ).json()
            if expect_replan:
                async with app.state.sessions() as session:
                    replan_rows = (
                        await session.scalars(
                            select(ReplanDecisionRow).where(
                                ReplanDecisionRow.workspace_id == workspace_id,
                                ReplanDecisionRow.task_id == UUID(task_id),
                            )
                        )
                    ).all()
                assert len(replan_rows) == 1
                assert replan_rows[0].status == "APPLIED"
                assert replan_rows[0].to_plan_id is not None
                assert replan_rows[0].payload["trigger"] == "EVIDENCE_INSUFFICIENT"
                current_steps = [
                    item
                    for item in steps
                    if item["plan_id"] == str(replan_rows[0].to_plan_id)
                ]
                assert len(steps) == 19
                assert len(current_steps) == 10
                assert {item["state"] for item in current_steps} == {"COMPLETED"}
            else:
                assert len(steps) == 9
                assert {item["state"] for item in steps} == {"COMPLETED"}
            artifacts = (
                await client.get(
                    "/artifacts",
                    headers=owner,
                    params={"workspace_id": str(workspace_id)},
                )
            ).json()
            assert len(artifacts) == 1
            if bound_hash is None:
                bound_hash = artifacts[0]["sha256"]
            assert artifacts[0]["sha256"] == bound_hash
            downloaded = await client.get(
                f"/artifacts/{artifacts[0]['artifact_id']}/download",
                headers=owner,
                params={"workspace_id": str(workspace_id)},
            )
            assert downloaded.status_code == 200
            assert hashlib.sha256(downloaded.content).hexdigest() == bound_hash
            report = Document(__import__("io").BytesIO(downloaded.content))
            report_text = "\n".join(item.text for item in report.paragraphs)
            for heading in (
                "Executive Summary",
                "Supplier Profile",
                "Risk Rating",
                "Risk Matrix",
                "Key Findings",
                "Evidence",
                "Missing Information",
                "Recommendation",
                "Citation",
            ):
                assert heading in report_text

            async with app.state.sessions() as session:
                claim_count = await session.scalar(
                    select(func.count()).select_from(ClaimRow).where(
                        ClaimRow.workspace_id == workspace_id,
                        ClaimRow.task_id == UUID(task_id),
                    )
                )
                replan_count = await session.scalar(
                    select(func.count()).select_from(ReplanDecisionRow).where(
                        ReplanDecisionRow.workspace_id == workspace_id,
                        ReplanDecisionRow.task_id == UUID(task_id),
                    )
                )
                artifact_count = await session.scalar(
                    select(func.count()).select_from(ArtifactRow).where(
                        ArtifactRow.workspace_id == workspace_id,
                        ArtifactRow.task_id == UUID(task_id),
                    )
                )
            assert claim_count == (3 if expect_replan else 2)
            assert replan_count == (1 if expect_replan else 0)
            assert artifact_count == 1
            trace_response = await client.get(
                f"/tasks/{task_id}/trace",
                headers=owner,
                params={"workspace_id": str(workspace_id)},
            )
            assert trace_response.status_code == 200
            trace = trace_response.json()
            trace_kinds = {item["kind"] for item in trace["nodes"]}
            assert {
                "TASK",
                "ROUTER",
                "PLAN",
                "STEP",
                "ATTEMPT",
                "EVIDENCE",
                "CAPABILITY",
                "LLM",
                "MCP",
                "ARTIFACT",
                "FINISHED",
            } <= trace_kinds
            if expect_approval:
                assert "HITL" in trace_kinds
            if expect_replan:
                assert "REPLAN" in trace_kinds
            node_ids = {item["node_id"] for item in trace["nodes"]}
            assert all(
                item["parent_id"] is None or item["parent_id"] in node_ids
                for item in trace["nodes"]
            )
            async with procurement() as session:
                admissions = (
                    await session.scalars(
                        select(AdmissionRow).where(
                            AdmissionRow.workspace_id == workspace_id
                        )
                    )
                ).all()
            assert len(admissions) == (1 if expect_approval else 0)
            if expect_approval:
                assert admissions[0].supplier_id == supplier.supplier_id
                assert admissions[0].report_sha256 == bound_hash
            assert FlagshipProvider.calls["/v1/chat/completions"] == (
                4 if expect_replan else 2
            )
            assert FlagshipProvider.gap_calls == (2 if expect_replan else 0)
            completed = True
    finally:
        if worker is not None and worker.poll() is None:
            worker.kill()
            worker.wait(timeout=10)
        if handle is not None and not completed:
            try:
                await handle.terminate("supplier flagship cleanup")
            except RPCError:
                pass
        await procurement_engine.dispose()
        provider.shutdown()
        provider.server_close()
        provider_thread.join(timeout=5)
