import hashlib
import json
import zipfile
from pathlib import Path
from typing import cast

from docx import Document
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from enterprise_agent.procurement.demo import DEMO_WORKSPACE_ID, demo_suppliers

ROOT = Path("demo/supplier-risk")
GENERATED = ROOT / "generated"
SKILL_SOURCE = Path("skills/supplier-risk")
SKILL_MEMBERS = (
    "skill.yaml",
    "instructions.md",
    "workflow.yaml",
    "schemas/input.schema.json",
    "schemas/output.schema.json",
    "policies/permissions.yaml",
    "eval/cases.json",
)
POLICY_LINES = (
    "Enterprise Procurement Admission Policy v1.0",
    "1. Supplier registration status must be ACTIVE.",
    "2. Any CRITICAL integrity or fraud risk in the last 24 months requires rejection.",
    "3. At least three completed orders and a 90 percent on-time delivery rate are required.",
    "4. An ACTIVE prior or trial contract is required for standard admission.",
    "5. Missing screening, history, or contract evidence requires additional research and replan.",
    "6. Admission creation is an external side effect and requires explicit human approval.",
)


def minimal_pdf(lines: tuple[str, ...]) -> bytes:
    escaped = [line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") for line in lines]
    stream = "BT /F1 11 Tf 50 790 Td 15 TL " + " T* ".join(
        f"({line}) Tj" for line in escaped
    ) + " ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        f"<< /Length {len(stream.encode())} >>\nstream\n{stream}\nendstream".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{index} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = len(data)
    data.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        data.extend(f"{offset:010d} 00000 n \n".encode())
    data.extend(
        f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(data)


def main() -> None:
    GENERATED.mkdir(parents=True, exist_ok=True)
    suppliers = demo_suppliers()

    document = Document()
    document.add_heading(POLICY_LINES[0], 0)
    for line in POLICY_LINES[1:]:
        document.add_paragraph(line)
    document.save(str(GENERATED / "enterprise-procurement-policy.docx"))
    (GENERATED / "enterprise-procurement-policy.pdf").write_bytes(minimal_pdf(POLICY_LINES))

    workbook = Workbook()
    summary = cast(Worksheet, workbook.active)
    summary.title = "Suppliers"
    summary.append(
        ["supplier_id", "label", "legal_name", "registration_number", "status", "expected_outcome"]
    )
    history = workbook.create_sheet("Purchase History")
    history.append(["supplier_id", "order_id", "amount", "currency", "delivery_status"])
    risks = workbook.create_sheet("Risk Events")
    risks.append(["supplier_id", "event_id", "severity", "description", "source"])
    contracts = workbook.create_sheet("Contracts")
    contracts.append(["supplier_id", "contract_id", "title", "status", "signed_at"])
    for supplier in suppliers:
        summary.append(
            [
                str(supplier.supplier_id),
                supplier.label,
                supplier.profile["legal_name"],
                supplier.profile["registration_number"],
                supplier.profile["status"],
                supplier.expected_outcome,
            ]
        )
        for item in supplier.purchase_history:
            history.append([str(supplier.supplier_id), *item.values()])
        for item in supplier.risk_flags:
            risks.append([str(supplier.supplier_id), *item.values()])
        for item in supplier.contracts:
            contracts.append([str(supplier.supplier_id), *item.values()])
    workbook.save(GENERATED / "supplier-evidence.xlsx")

    manifest = {
        "workspace_id": str(DEMO_WORKSPACE_ID),
        "suppliers": [
            {
                "supplier_id": str(item.supplier_id),
                "label": item.label,
                "expected_outcome": item.expected_outcome,
                "expected_reason": item.expected_reason,
            }
            for item in suppliers
        ],
    }
    manifest_path = GENERATED / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    with zipfile.ZipFile(
        GENERATED / "supplier-risk.skill.zip",
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as archive:
        for name in SKILL_MEMBERS:
            info = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o100644 << 16
            archive.writestr(info, (SKILL_SOURCE / name).read_bytes())
    hashes = {
        path.name: hashlib.sha256(path.read_bytes()).hexdigest()
        for path in sorted(GENERATED.iterdir())
        if path.name != "SHA256SUMS.json"
    }
    (GENERATED / "SHA256SUMS.json").write_text(
        json.dumps(hashes, indent=2) + "\n", encoding="utf-8"
    )


if __name__ == "__main__":
    main()
